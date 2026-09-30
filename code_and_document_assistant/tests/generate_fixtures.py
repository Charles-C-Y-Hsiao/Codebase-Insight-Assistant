"""Generate synthetic binary documents for parser tests (no company information)."""
import sys
import zipfile
from pathlib import Path
from datetime import datetime
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from openpyxl import Workbook
from reportlab.pdfgen import canvas

target = Path(sys.argv[1])
target.mkdir(parents=True, exist_ok=True)
doc = Document()
doc.add_heading('測試政策', level=1)
doc.add_paragraph('產品甲，2026 第一季，單位 TWD。')
table = doc.add_table(rows=2, cols=2)
table.cell(0, 0).text, table.cell(0, 1).text = '產品', '金額 TWD'
table.cell(1, 0).text, table.cell(1, 1).text = '甲', '1234'
doc.save(target / 'policy.docx')
deck = Presentation()
for title in ('First slide', 'Second slide'):
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = title
    slide.placeholders[1].text = 'Quarter 2026 Q1'
table = deck.slides[1].shapes.add_table(2,2, Inches(1),Inches(3), Inches(6),Inches(1)).table
table.cell(0,0).text,table.cell(0,1).text='Product','Amount TWD'
table.cell(1,0).text,table.cell(1,1).text='A','1234'
deck.save(target / 'slides.pptx')
wb = Workbook()
ws = wb.active
ws.title = '銷售'
ws.append(['產品','金額 TWD','日期','計算'])
ws.append(['甲',1234,datetime(2026,1,2),'=SUM(B2:B2)'])
ws.merge_cells('E1:F1')
ws['E1']='合併標題'
wb.save(target / 'sales.xlsx')
c = canvas.Canvas(str(target / 'sample.pdf'))
c.drawString(72,720,'Orders 2026 Q1, Product A, TWD 1234')
c.showPage()
c.drawString(72,720,'Page two evidence')
c.showPage()
c.save()
c = canvas.Canvas(str(target / 'blank.pdf'))
c.showPage()
c.save()
with zipfile.ZipFile(target / 'entity.docx', 'w') as archive:
    archive.writestr('word/document.xml', '<!DOCTYPE x [<!ENTITY secret SYSTEM "file:///not-authorized">]><x>&secret;</x>')
print('Synthetic PDF / DOCX / PPTX / XLSX fixtures generated')
